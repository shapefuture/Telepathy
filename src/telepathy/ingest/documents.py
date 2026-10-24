"""Document and transcript ingestion (P07-01..P07-05, P07-09).

Files posted in an enabled group inherit that group's visibility and its message link
(deleting the message revokes the document). Binaries are not kept: only a SHA-256,
the parser version and the extracted excerpts. Scanned PDFs are marked `needs_ocr`;
nothing is sent to a third-party OCR service.
"""

from __future__ import annotations

import hashlib
import io
import re
import sqlite3
from dataclasses import dataclass, field
from html.parser import HTMLParser

from telepathy.domain.text import normalize, sha256_hex
from telepathy.domain.types import new_id, to_iso, utcnow
from telepathy.storage.db import transaction

PARSER_VERSION = "1"
MAX_BYTES = 20 * 1024 * 1024  # Bot API download limit
SUPPORTED = {".txt", ".md", ".markdown", ".html", ".htm", ".vtt", ".srt", ".pdf", ".docx", ".pptx"}

# Chunking (P05-03): ~350 tokens target, 600 max, 70 overlap. Approximated in words.
TARGET_WORDS = 260
MAX_WORDS = 450
OVERLAP_WORDS = 50
TRANSCRIPT_WINDOW_SECONDS = 90


@dataclass
class Segment:
    locator: str | None
    text: str


@dataclass
class ParsedDocument:
    status: str  # parsed | needs_ocr | unsupported | failed
    fmt: str
    segments: list[Segment] = field(default_factory=list)
    error: str | None = None


def extension(file_name: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", file_name or "")
    return m.group(1).lower() if m else ""


def is_supported(file_name: str | None, size: int | None) -> bool:
    return (
        bool(file_name)
        and extension(file_name or "") in SUPPORTED
        and (size is None or size <= MAX_BYTES)
    )


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


_TS = re.compile(r"(?:(\d{1,2}):)?(\d{1,2}):(\d{2})[.,]\d{1,3}\s*-->")


def _fmt_ts(seconds: int) -> str:
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def parse_transcript(text: str) -> list[Segment]:
    """VTT/SRT -> time-window segments with mm:ss locators (P07-01, P07-05)."""
    cues: list[tuple[int, str]] = []
    current_start: int | None = None
    buf: list[str] = []
    for line in text.splitlines():
        m = _TS.search(line)
        if m:
            if current_start is not None and buf:
                cues.append((current_start, " ".join(buf)))
            h, mi, s = m.group(1), m.group(2), m.group(3)
            current_start = int(h or 0) * 3600 + int(mi) * 60 + int(s)
            buf = []
        elif (
            line.strip()
            and not line.strip().isdigit()
            and line.strip() != "WEBVTT"
            and current_start is not None
            and not line.startswith("NOTE")
        ):
            buf.append(re.sub(r"<[^>]+>", "", line.strip()))
    if current_start is not None and buf:
        cues.append((current_start, " ".join(buf)))
    segments: list[Segment] = []
    win_start: int | None = None
    win: list[str] = []
    for start, cue in cues:
        if win_start is None:
            win_start = start
        if start - win_start >= TRANSCRIPT_WINDOW_SECONDS and win:
            segments.append(Segment(_fmt_ts(win_start), " ".join(win)))
            win_start, win = start, []
        win.append(cue)
    if win and win_start is not None:
        segments.append(Segment(_fmt_ts(win_start), " ".join(win)))
    return segments


def parse(file_name: str, data: bytes) -> ParsedDocument:
    ext = extension(file_name)
    fmt = ext.lstrip(".")
    if ext not in SUPPORTED:
        return ParsedDocument("unsupported", fmt)
    try:
        if ext in (".txt", ".md", ".markdown"):
            return ParsedDocument("parsed", fmt, [Segment(None, data.decode("utf-8", "replace"))])
        if ext in (".html", ".htm"):
            p = _TextExtractor()
            p.feed(data.decode("utf-8", "replace"))
            return ParsedDocument("parsed", fmt, [Segment(None, "".join(p.parts))])
        if ext in (".vtt", ".srt"):
            return ParsedDocument(
                "parsed", fmt, parse_transcript(data.decode("utf-8-sig", "replace"))
            )
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = [
                Segment(f"p. {i}", page.extract_text() or "")
                for i, page in enumerate(reader.pages, 1)
            ]
            if not any(len(s.text.strip()) > 20 for s in pages):
                return ParsedDocument("needs_ocr", fmt)  # P07-03: no hidden OCR submission
            return ParsedDocument("parsed", fmt, pages)
        if ext == ".docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in d.paragraphs)
            return ParsedDocument("parsed", fmt, [Segment(None, text)])
        if ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
                slides.append(Segment(f"slide {i}", "\n".join(texts)))
            return ParsedDocument("parsed", fmt, slides)
    except Exception as exc:  # malformed files must not crash the worker
        return ParsedDocument("failed", fmt, error=f"{type(exc).__name__}: {exc}"[:300])
    return ParsedDocument("unsupported", fmt)


def chunk(segments: list[Segment]) -> list[Segment]:
    """Split long segments into overlapping word windows; keep each segment's locator."""
    out: list[Segment] = []
    for seg in segments:
        words = seg.text.split()
        if not words:
            continue
        if len(words) <= MAX_WORDS:
            out.append(Segment(seg.locator, " ".join(words)))
            continue
        start = 0
        while start < len(words):
            end = min(len(words), start + TARGET_WORDS)
            out.append(Segment(seg.locator, " ".join(words[start:end])))
            if end == len(words):
                break
            start = end - OVERLAP_WORDS
    return out


def store_document(
    conn: sqlite3.Connection,
    *,
    source_id: str,
    message_row_id: str,
    citation_uri: str | None,
    event_at: str,
    file_name: str,
    data: bytes,
) -> tuple[str, str, int]:
    """Parse + persist. Returns (document_id, status, excerpt_count). Idempotent per message."""
    digest = hashlib.sha256(data).hexdigest()
    parsed = parse(file_name, data)
    chunks = chunk(parsed.segments) if parsed.status == "parsed" else []
    now = to_iso(utcnow())
    with transaction(conn):
        existing = conn.execute(
            "SELECT id, status, excerpt_count FROM documents WHERE message_id=? AND sha256=?",
            (message_row_id, digest),
        ).fetchone()
        if existing:
            return existing["id"], existing["status"], existing["excerpt_count"]
        live = conn.execute(
            "SELECT 1 FROM messages WHERE id=? AND removed_at IS NULL", (message_row_id,)
        ).fetchone()
        if live is None:
            return "", "removed", 0
        doc_id = new_id()
        conn.execute(
            "INSERT INTO documents(id, source_id, path, sha256, format, parser_version, "
            "created_at, message_id, title, status, excerpt_count) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (
                doc_id,
                source_id,
                file_name,
                digest,
                parsed.fmt,
                PARSER_VERSION,
                now,
                message_row_id,
                file_name,
                parsed.status,
                len(chunks),
            ),
        )
        for c in chunks:
            body = f"[{file_name}{' · ' + c.locator if c.locator else ''}] {c.text}"
            cur = conn.execute(
                "INSERT INTO excerpts(id, source_id, message_id, document_id, body, citation_uri, "
                "event_at, valid_from, content_sha256, locator) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (
                    new_id(),
                    source_id,
                    message_row_id,
                    doc_id,
                    body,
                    citation_uri,
                    event_at,
                    now,
                    sha256_hex(body),
                    c.locator,
                ),
            )
            conn.execute(
                "INSERT INTO excerpts_fts(rowid, body) VALUES (?, ?)",
                (cur.lastrowid, normalize(body)),
            )
    return doc_id, parsed.status, len(chunks)
