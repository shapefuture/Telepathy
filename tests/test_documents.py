from __future__ import annotations

import io

import pytest
from conftest import ctx_for, enable, post

from telepathy.ingest import documents
from telepathy.ingest.service import message_row_id, remove_message
from telepathy.retrieval.answer import AnswerStatus, answer

VTT = """WEBVTT

1
00:00:01.000 --> 00:00:05.000
Welcome to the meetup.

2
00:02:00.000 --> 00:02:05.000
The <b>parking</b> is behind the library.
"""


def test_transcript_segments_have_timestamps():
    segs = documents.parse_transcript(VTT)
    assert [s.locator for s in segs] == ["00:01", "02:00"]
    assert "parking is behind" in segs[1].text


def test_chunking_overlaps_and_bounds():
    long = documents.Segment("p. 1", " ".join(f"w{i}" for i in range(1000)))
    chunks = documents.chunk([long])
    assert all(len(c.text.split()) <= documents.MAX_WORDS for c in chunks)
    assert (
        chunks[0].text.split()[-documents.OVERLAP_WORDS :]
        == chunks[1].text.split()[: documents.OVERLAP_WORDS]
    )


def _docx_bytes(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _blank_pdf() -> bytes:
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(200, 200)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    "name,data,status",
    [
        ("a.txt", b"plain text body", "parsed"),
        ("a.md", b"# Title\nbody", "parsed"),
        ("a.html", b"<p>hi</p><script>evil()</script>", "parsed"),
        ("a.docx", None, "parsed"),
        ("a.pptx", None, "parsed"),
        ("scan.pdf", None, "needs_ocr"),
        ("a.exe", b"MZ", "unsupported"),
        ("broken.docx", b"not a zip", "failed"),
    ],
)
def test_parse_formats(name, data, status):
    if data is None:
        data = {
            "a.docx": _docx_bytes("budget plan"),
            "a.pptx": _pptx_bytes("roadmap slide"),
            "scan.pdf": _blank_pdf(),
        }[name]
    parsed = documents.parse(name, data)
    assert parsed.status == status
    if name == "a.html":
        assert "evil" not in parsed.segments[0].text
    if name == "a.pptx":
        assert parsed.segments[0].locator == "slide 1"


def test_document_is_searchable_and_revoked_with_its_message(conn, owner):
    sid = enable(conn, owner, -1001234)
    post(conn, -1001234, "slides attached", user_id=1, message_id=50)
    row = message_row_id(conn, sid, "50")
    doc_id, status, n = documents.store_document(
        conn,
        source_id=sid,
        message_row_id=row,
        citation_uri="https://t.me/c/1234/50",
        event_at="2026-10-01T00:00:00+00:00",
        file_name="meetup.vtt",
        data=VTT.encode(),
    )
    assert status == "parsed" and n == 2
    # Idempotent per message + content hash.
    assert (
        documents.store_document(
            conn,
            source_id=sid,
            message_row_id=row,
            citation_uri=None,
            event_at="2026-10-01T00:00:00+00:00",
            file_name="meetup.vtt",
            data=VTT.encode(),
        )[0]
        == doc_id
    )
    ans = answer(conn, ctx_for(conn, 1), "where is the parking?")
    assert ans.status is AnswerStatus.ANSWERED
    assert "02:00" in ans.quotes[0].text
    remove_message(conn, sid, "50", actor_member_id=owner, reason="test")
    assert (
        answer(conn, ctx_for(conn, 1), "where is the parking?").status is AnswerStatus.INSUFFICIENT
    )
    assert conn.execute("SELECT status FROM documents").fetchone()[0] == "removed"


def test_caption_edit_keeps_document_excerpts(conn, owner):
    sid = enable(conn, owner, -1001)
    post(conn, -1001, "old caption", user_id=1, message_id=60)
    row = message_row_id(conn, sid, "60")
    documents.store_document(
        conn,
        source_id=sid,
        message_row_id=row,
        citation_uri=None,
        event_at="2026-10-01T00:00:00+00:00",
        file_name="n.txt",
        data=b"the wifi password is on the fridge",
    )
    post(conn, -1001, "new caption", user_id=1, message_id=60, edit_date=1_760_000_900)
    live = [r[0] for r in conn.execute("SELECT body FROM excerpts WHERE valid_to IS NULL")]
    assert "new caption" in live and any("wifi" in b for b in live)


def test_is_supported():
    assert documents.is_supported("x.PDF", 1000)
    assert not documents.is_supported("x.pdf", documents.MAX_BYTES + 1)
    assert not documents.is_supported(None, 10)
